"""Dump the semantic content of every slide, splitting template chrome from content."""
import sys, json
from pptx import Presentation
from pptx.util import Emu

# shapes that belong to the shared page chrome (all decks share these positions)
CHROME_POS = {
    (0.00, 0.00), (0.40, -0.04), (0.55, -0.26), (0.61, 0.94),
    (11.54, 5.71), (12.36, 3.35), (12.84, 0.86), (0.00, 0.0),
    (9.36, 0.60), (12.37, 0.60), (9.80, 0.76), (9.83, 0.77),
}
CHROME_TEXT = {'哈理工A.I.R.创新实验室', '哈理工 A.I.R. 创新实验室',
               '『生无所息，斗无所止』'}


def is_chrome(sh):
    L = round(Emu(sh.left).inches, 2)
    T = round(Emu(sh.top).inches, 2)
    if (L, T) in {(0.0, 0.0), (0.4, -0.04), (0.55, -0.26), (0.61, 0.94),
                  (11.54, 5.71), (12.36, 3.35), (12.84, 0.86),
                  (9.36, 0.6), (12.37, 0.6), (9.8, 0.76), (9.83, 0.77)}:
        return True
    if sh.has_text_frame and sh.text_frame.text.strip() in CHROME_TEXT:
        return True
    return False


def dump(path):
    prs = Presentation(path)
    out = []
    for i, s in enumerate(prs.slides, 1):
        content = []
        for sh in s.shapes:
            if is_chrome(sh):
                continue
            L, T = round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2)
            W, H = round(Emu(sh.width).inches, 2), round(Emu(sh.height).inches, 2)
            item = {'t': str(sh.shape_type), 'pos': [L, T, W, H], 'name': sh.name}
            if sh.has_text_frame:
                paras = []
                for p in sh.text_frame.paragraphs:
                    runs = [(r.text, r.font.size.pt if r.font.size else None,
                             r.font.bold, r.font.name,
                             str(r.font.color.rgb) if (r.font.color and r.font.color.type == 1) else None)
                            for r in p.runs]
                    if any(r[0].strip() for r in runs):
                        paras.append(runs)
                item['paras'] = paras
                item['text'] = '\n'.join(''.join(r[0] for r in pa) for pa in paras)
                item['align'] = str(sh.text_frame.paragraphs[0].alignment)
            if sh.has_table:
                item['table'] = [[c.text for c in row.cells] for row in sh.table.rows]
            if sh.shape_type == 13:
                item['pic'] = True
            content.append(item)
        out.append({'n': i, 'layout': s.slide_layout.name, 'content': content})
    return out


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    for p in sys.argv[1:]:
        print('########## ' + p)
        print(json.dumps(dump(p), ensure_ascii=False, indent=0))
