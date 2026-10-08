"""Classify every non-chrome shape into a layout block kind, per slide."""
import sys
from pptx import Presentation
from pptx.util import Emu

sys.stdout.reconfigure(encoding='utf-8', errors='replace')
CHROME_TEXT = {'哈理工A.I.R.创新实验室', '哈理工 A.I.R. 创新实验室', '『生无所息，斗无所止』'}
CHROME_POS = {(0.0, 0.0), (0.4, -0.04), (0.55, -0.26), (0.61, 0.94),
              (11.54, 5.71), (12.36, 3.35), (12.84, 0.86),
              (9.36, 0.6), (12.37, 0.6), (9.8, 0.76), (9.83, 0.77)}


def geom(sh):
    return (round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2),
            round(Emu(sh.width).inches, 2), round(Emu(sh.height).inches, 2))


def is_chrome(sh):
    if sh.left is None:
        return False
    L, T, W, H = geom(sh)
    if (L, T) in CHROME_POS:
        return True
    if sh.has_text_frame and sh.text_frame.text.strip() in CHROME_TEXT:
        return True
    return False


def kind(sh):
    L, T, W, H = geom(sh)
    txt = sh.text_frame.text.strip() if sh.has_text_frame else ''
    mono = any(r.font.name in ('Consolas', 'DejaVu Sans Mono', 'Courier New')
               for p in sh.text_frame.paragraphs for r in p.runs) if sh.has_text_frame else False
    if sh.has_table:
        return 'NATIVE_TABLE(%dx%d)' % (len(sh.table.rows), len(sh.table.columns))
    if sh.shape_type == 13:
        return 'PICTURE(%.2f,%.2f,%.2fx%.2f)' % (L, T, W, H)
    if abs(L - 1.54) < 0.03 and abs(T - 0.46) < 0.03:
        return 'TITLE'
    if abs(L - 0.87) < 0.02 and T >= 5.0 and W > 11:
        return 'BAR_FILL'
    if abs(L - 0.87) < 0.02 and T >= 5.0 and W > 11:
        return 'BAR_TEXT'
    if abs(L - 0.87) < 0.02 and T >= 5.0:
        return 'BAR?'
    if abs(L - 0.87) < 0.02 and abs(W - 0.78) < 0.05:
        return 'HWNUM(%s)' % txt[:2]
    if abs(L - 1.72) < 0.03 and W > 10:
        return 'HWTEXT(%.2f,%.2f)' % (T, H)
    if abs(L - 0.87) < 0.02 and abs(W - 5.97) < 0.05:
        return 'ITEM(%.2f)' % T
    if abs(L - 7.08) < 0.02 and abs(W - 5.25) < 0.05:
        return 'CARD'
    if abs(L - 7.25) < 0.02:
        return 'CARD_CODE(%.2f)' % T
    if abs(L - 0.87) < 0.02 and abs(W - 11.50) < 0.05 and H < 0.5:
        return 'TBL_HEADER'
    if abs(L - 0.87) < 0.02 and abs(W - 11.50) < 0.05:
        return 'TBL_BAND'
    if abs(W - 11.36) < 0.05 and T > 5.0:
        return 'BAR_TEXT'
    if T > 5.0 and W > 11:
        return 'BOTTOM_TEXT(%.2f)' % T
    if mono:
        return 'CODE(%.2f,%.2f,%.2fx%.2f)' % (L, T, W, H)
    return 'TEXT(%.2f,%.2f,%.2fx%.2f)' % (L, T, W, H)


for path in sys.argv[1:]:
    prs = Presentation(path)
    print('##### ' + path.split('\\')[-1])
    for i, s in enumerate(prs.slides, 1):
        ks = [kind(sh) for sh in s.shapes if not is_chrome(sh)]
        print('p%-3d %s' % (i, '  '.join(ks)))
