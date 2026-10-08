"""Comprehensive style fingerprint for each deck: fonts, colors, fills, layout patterns."""
import sys, json, zipfile, collections
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.util import Emu

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def q(tag):
    return '{%s}%s' % (A, tag)


def collect(path):
    prs = Presentation(path)
    fonts = collections.Counter()
    sizes = collections.Counter()
    textcolors = collections.Counter()
    fills = collections.Counter()
    lines = collections.Counter()
    mono = collections.Counter()
    titles = []
    bottom_bar = 0
    badges = 0
    ntab = 0
    pics = 0
    per_slide = []

    def rgb(v):
        return str(v) if v is not None else None

    for si, s in enumerate(prs.slides, 1):
        info = {'n': si, 'title': None, 'title_pos': None}
        has_bar = False
        has_badge = False
        for sh in s.shapes:
            stack = [sh]
            while stack:
                cur = stack.pop()
                if cur.shape_type == 6:
                    stack.extend(list(cur.shapes))
                    continue
                if cur.shape_type == 19:
                    ntab += 1
                if cur.shape_type == 13:
                    pics += 1
                # fill / line
                try:
                    f = cur.fill
                    if f.type is not None and str(f.type) == 'MSO_FILL_TYPE.SOLID (1)':
                        fills[rgb(f.fore_color.rgb) if f.fore_color.type == 1 else 'theme'] += 1
                except Exception:
                    pass
                try:
                    if cur.line.fill.type is not None and cur.line.color and cur.line.color.type == 1:
                        lines[rgb(cur.line.color.rgb)] += 1
                except Exception:
                    pass
                if not cur.has_text_frame:
                    continue
                tf = cur.text_frame
                for p in tf.paragraphs:
                    txt = ''.join(r.text for r in p.runs)
                    for r in p.runs:
                        fnt = r.font
                        name = fnt.name
                        if name:
                            fonts[name] += 1
                        if fnt.size:
                            sizes[round(fnt.size.pt, 2)] += 1
                        try:
                            if fnt.color and fnt.color.type == 1:
                                textcolors[rgb(fnt.color.rgb)] += 1
                        except Exception:
                            pass
                        if name and name in ('Consolas', 'DejaVu Sans Mono', 'Courier New'):
                            mono[name] += 1
                    if txt.strip():
                        t = txt.strip()
                        if t.startswith('本页判断'):
                            has_bar = True
                        if t in ('1.', '2.', '3.', '4.'):
                            has_badge = True
                        if len(t) < 40 and info['title'] is None and len(t) > 3 and 'A.I.R' not in t and not t.isdigit():
                            info['title'] = t
                            info['title_pos'] = (round(Emu(cur.left).inches, 2), round(Emu(cur.top).inches, 2),
                                                 round(Emu(cur.width).inches, 2), round(Emu(cur.height).inches, 2))
        if has_bar:
            bottom_bar += 1
        if has_badge:
            badges += 1
        per_slide.append(info)
    return {
        'path': path,
        'slides': len(prs.slides),
        'fonts': dict(fonts.most_common()),
        'fontSizes': dict(sizes.most_common(15)),
        'textColors': dict(textcolors.most_common(20)),
        'solidFills': dict(fills.most_common(15)),
        'lineColors': dict(lines.most_common(10)),
        'tables': ntab,
        'pictures': pics,
        'slidesWithBottomBar': bottom_bar,
        'slidesWithNumberBadges': badges,
        'slideTitles': per_slide,
    }


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    for p in sys.argv[1:]:
        print(json.dumps(collect(p), ensure_ascii=False, indent=1))
        print()
