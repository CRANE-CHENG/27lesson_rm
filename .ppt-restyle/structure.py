"""Compact per-slide structural listing of content shapes."""
import sys
from pptx import Presentation
from pptx.util import Emu

CHROME_TEXT = {'哈理工A.I.R.创新实验室', '哈理工 A.I.R. 创新实验室', '『生无所息，斗无所止』'}
CHROME_POS = {(0.0, 0.0), (0.4, -0.04), (0.55, -0.26), (0.61, 0.94),
              (11.54, 5.71), (12.36, 3.35), (12.84, 0.86),
              (9.36, 0.6), (12.37, 0.6), (9.8, 0.76), (9.83, 0.77)}


def is_chrome(sh):
    if sh.left is None:
        return False
    L, T = round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2)
    if (L, T) in CHROME_POS:
        return True
    if sh.has_text_frame and sh.text_frame.text.strip() in CHROME_TEXT:
        return True
    return False


def first(sh):
    if sh.has_text_frame:
        for p in sh.text_frame.paragraphs:
            t = ''.join(r.text for r in p.runs).strip()
            if t:
                return t.replace('\n', ' / ')[:46]
    if sh.has_table:
        return 'TABLE %dx%d :: %s' % (len(sh.table.rows), len(sh.table.columns),
                                      ' | '.join(c.text.replace('\n', ' ')[:12]
                                                 for c in sh.table.rows[0].cells))
    return ''


try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

for path in sys.argv[1:]:
    prs = Presentation(path)
    print('##### ' + path.split('\\')[-1])
    for i, s in enumerate(prs.slides, 1):
        rows = []
        for sh in s.shapes:
            if is_chrome(sh):
                continue
            L, T = round(Emu(sh.left).inches, 2), round(Emu(sh.top).inches, 2)
            W, H = round(Emu(sh.width).inches, 2), round(Emu(sh.height).inches, 2)
            kind = {5: 'free', 1: 'shape', 17: 'text', 13: 'PIC', 19: 'TBL'}.get(sh.shape_type, str(sh.shape_type))
            rows.append('   %-5s (%5.2f,%5.2f) %5.2fx%5.2f  %s' % (kind, L, T, W, H, first(sh)))
        if rows:
            print('-- p%d [%s]' % (i, s.slide_layout.name or '-'))
            print('\n'.join(rows))
        else:
            print('-- p%d [%s]  (no content shapes = 章节/封面/目录/结尾页)' % (i, s.slide_layout.name or '-'))
