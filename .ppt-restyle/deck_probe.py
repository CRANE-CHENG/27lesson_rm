"""Dump PPTX theme + slide structure for style comparison."""
import sys, json, zipfile
from xml.etree import ElementTree as ET
from pptx import Presentation
from pptx.util import Emu

A = 'http://schemas.openxmlformats.org/drawingml/2006/main'
P = 'http://schemas.openxmlformats.org/presentationml/2006/main'


def q(ns, tag):
    return '{%s}%s' % (ns, tag)


def dump_theme(path):
    out = {'path': path}
    with zipfile.ZipFile(path) as z:
        names = z.namelist()
        out['parts'] = len(names)
        out['masters'] = sorted(n for n in names if n.startswith('ppt/slideMasters/') and n.endswith('.xml'))
        out['layouts'] = sorted(n for n in names if n.startswith('ppt/slideLayouts/') and n.endswith('.xml'))
        out['media'] = sorted(n for n in names if n.startswith('ppt/media/'))
        # presentation.xml
        pr = ET.fromstring(z.read('ppt/presentation.xml'))
        sz = pr.find(q(P, 'sldSz'))
        out['sldSz'] = (int(sz.get('cx')), int(sz.get('cy')))
        # theme1
        th = ET.fromstring(z.read('ppt/theme/theme1.xml'))
        cs = th.find(q(A, 'themeElements')).find(q(A, 'clrScheme'))
        colors = {}
        for c in cs:
            tag = c.tag.split('}')[1]
            child = list(c)[0]
            colors[tag] = child.get('val') or child.get('lastClr')
        out['colorScheme'] = colors
        fs = th.find(q(A, 'themeElements')).find(q(A, 'fontScheme'))
        fonts = {}
        for kind in ('majorFont', 'minorFont'):
            el = fs.find(q(A, kind))
            fonts[kind] = {
                'latin': el.find(q(A, 'latin')).get('typeface'),
                'ea': el.find(q(A, 'ea')).get('typeface') if el.find(q(A, 'ea')) is not None else None,
                'cs': el.find(q(A, 'cs')).get('typeface') if el.find(q(A, 'cs')) is not None else None,
            }
        out['fontScheme'] = fonts
        out['themeName'] = th.get('name')
    # masters / layouts text
    prs = Presentation(path)
    out['nSlides'] = len(prs.slides)
    lay = []
    for i, l in enumerate(prs.slide_masters[0].slide_layouts):
        lay.append('%d:%s' % (i, l.name))
    out['masterLayouts'] = lay
    return out


def walk(shapes, indent, lines, depth=0):
    for sh in shapes:
        try:
            st = sh.shape_type
        except Exception:
            st = '?'
        info = '%s- [%s] %s name=%r pos=(%.2f,%.2f) size=(%.2f,%.2f)' % (
            '  ' * indent, st, sh.shape_id, sh.name,
            Emu(sh.left).inches if sh.left is not None else -1,
            Emu(sh.top).inches if sh.top is not None else -1,
            Emu(sh.width).inches if sh.width is not None else -1,
            Emu(sh.height).inches if sh.height is not None else -1)
        lines.append(info)
        if sh.shape_type == 6:  # group
            walk(sh.shapes, indent + 1, lines)
            continue
        if sh.has_text_frame:
            for p in sh.text_frame.paragraphs:
                txt = ''.join(r.text for r in p.runs)
                if not txt.strip():
                    continue
                r0 = p.runs[0]
                f = r0.font
                col = None
                try:
                    if f.color and f.color.type is not None:
                        col = str(f.color.rgb) if f.color.type == 1 else 'theme:%s' % f.color.theme_color
                except Exception:
                    col = 'err'
                lines.append('%s    P lvl=%s align=%s sz=%s bold=%s font=%s color=%s | %s' % (
                    '  ' * indent, p.level,
                    p.alignment, f.size.pt if f.size else None, f.bold, f.name, col, txt[:70]))


def dump_slides(path, maxslides=999):
    prs = Presentation(path)
    lines = []
    for i, s in enumerate(prs.slides, 1):
        if i > maxslides:
            break
        lines.append('===== SLIDE %d  layout=%r =====' % (i, s.slide_layout.name))
        walk(s.shapes, 0, lines)
    return lines


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    mode = sys.argv[1]
    if mode == 'theme':
        for p in sys.argv[2:]:
            print(json.dumps(dump_theme(p), ensure_ascii=False, indent=1))
    elif mode == 'slides':
        n = 999
        paths = sys.argv[2:]
        if paths[0].isdigit():
            n = int(paths[0])
            paths = paths[1:]
        for p in paths:
            print('########## ' + p)
            print('\n'.join(dump_slides(p, n)))
