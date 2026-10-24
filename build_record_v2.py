# -*- coding: utf-8 -*-
"""重建交作业记录表：备注并入对应那次作业的“交作业情况”单元格，删除单独的备注列。"""
import shutil
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SRC = r"D:\桌面\培训\交作业情况记录.xlsx"
DST = r"D:\桌面\培训\交作业情况记录_已更新.xlsx"

# 姓名 -> (第一次作业单元格文本, 第二次作业单元格文本)
DATA = {
 "王玉琳": ("已交", "已交"),
 "顾柳霖": ("已交", "已交"),
 "刘永鹏": ("已交", "已交"),
 "唐俊毅": ("已交", "已交"),
 "郑瀚":   ("已交\n（2 次提交，以第二次为准）", "已交"),
 "刘岳":   ("已交", "未交"),
 "郝欣如": ("已交", "已交\n（缺源代码，仅截图与视频）"),
 "刘博宇": ("已交\n（缺演示视频）", "已交"),
 "熊磊":   ("未交\n（缺第一次作业）", "已交"),
 "蔡一炜": ("已交", "未交"),
 "黄浡桓": ("已交", "未交"),
 "殷浩然": ("已交", "未交"),
 "任宏杰": ("已交", "未交"),
 "黄竟轩": ("已交", "已交"),
 "周泽宇": ("已交", "未交"),
 "王子豪": ("已交", "未交"),
 "左翔宇": ("已交\n（2 次提交，以第二次为准）", "未交"),
 "王璟":   ("已交", "未交"),
 "田晓宁": ("已交", "已交"),
 "刘明远": ("已交\n（2 次提交，以第二次为准）", "未交"),
 "常凯":   ("已交", "未交"),
 "史昊田": ("已交", "已交"),
 "李韩煦": ("已交", "已交"),
 "李壮":   ("已交", "未交"),
 "姚硕":   ("已交", "未交"),
 "曹奥华": ("已交\n（4 次提交，以最后一次为准）", "已交"),
 "张博":   ("已交\n（源码以 .doc 提交）", "未交"),
 "段少龙": ("已交\n（缺演示视频）", "未交"),
 "宋文松": ("已交\n（3 次提交，以第三次为准）", "未交"),
 "何青铄": ("未交\n（仅提交无效快捷方式）", "未交"),
 "刘飞扬": ("已交\n（单文件合并完成三题）", "已交"),
 "李想":   ("已交\n（缺源代码，仅演示视频）", "未交"),
 "陈梦祺": ("已交", "已交"),
 "王鹏宇": ("已交", "未交"),
 "李泽冲": ("已交", "未交"),
 "程义硕": ("已交\n（2 次提交，第二次补齐代码）", "未交"),
 "孟繁志": ("已交", "未交"),
}
FIX = {"孟繁志": ("电气", "26-3"), "熊磊": (None, "26-2")}

shutil.copyfile(SRC, DST)
wb = openpyxl.load_workbook(DST)
ws = wb.active
ws._images = []                      # 去掉盖住 A1 的水印

FONT = "微软雅黑"
HEAD_FILL = PatternFill("solid", fgColor="FF17365D")
HEAD_FONT = Font(name=FONT, size=10, bold=True, color="FFFFFFFF")
GREEN_FILL = PatternFill("solid", fgColor="FFE2EFDA")
GREEN_FONT = Font(name=FONT, size=10.5, bold=True, color="FF375623")
RED_FILL = PatternFill("solid", fgColor="FFFCE4E4")
RED_FONT = Font(name=FONT, size=10.5, bold=True, color="FFC00000")
side = Side(style="thin", color="FFB8C7D9")
BORDER = Border(left=side, right=side, top=side, bottom=side)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

# 删除单独的备注列（I 列）
if ws.max_column >= 9:
    ws.delete_cols(9)

# 表头补齐到 8 列
ws.cell(1, 8, "第二次作业（作业分数）")
for col in range(1, 9):
    c = ws.cell(1, col)
    c.font = HEAD_FONT
    c.fill = HEAD_FILL
    c.alignment = CENTER
    c.border = BORDER
ws.row_dimensions[1].height = 38

last = 0
for r in range(2, ws.max_row + 1):
    name = ws.cell(r, 4).value
    if not name or name not in DATA:
        continue
    last = r
    if name in FIX:
        mj, cl = FIX[name]
        if mj:
            ws.cell(r, 2, mj)
        if cl:
            ws.cell(r, 3, cl)
    ws.row_dimensions[r].height = 32
    for col, text in ((5, DATA[name][0]), (7, DATA[name][1])):
        c = ws.cell(r, col, text)
        c.alignment = CENTER
        c.border = BORDER
        if text.startswith("已交"):
            c.fill, c.font = GREEN_FILL, GREEN_FONT
        else:
            c.fill, c.font = RED_FILL, RED_FONT
    for col in (1, 2, 3, 4, 6, 8):
        c = ws.cell(r, col)
        c.border = BORDER
        c.alignment = CENTER
    ws.cell(r, 6).value = None
    ws.cell(r, 8).value = None

for k, v in {"A": 5.0, "B": 10.6, "C": 8.4, "D": 8.4,
             "E": 30.0, "F": 24.0, "G": 30.0, "H": 24.0}.items():
    ws.column_dimensions[k].width = v
ws.freeze_panes = "A2"

# 统计区
n1 = sum(1 for v in DATA.values() if v[0].startswith("已交"))
n2 = sum(1 for v in DATA.values() if v[1].startswith("已交"))
total = len(DATA)
s = last + 2
for i, t in enumerate(["统计项", "应交人数", "已交人数", "未交人数", "提交率"]):
    col = 1 if i == 0 else i + 2
    c = ws.cell(s, col, t)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFFFF")
    c.fill = PatternFill("solid", fgColor="FF2E5C8A")
    c.alignment = CENTER
    c.border = BORDER
ws.merge_cells(start_row=s, start_column=1, end_row=s, end_column=2)
ws.row_dimensions[s].height = 24

for j, row in enumerate([("第一次作业", total, n1, total - n1, "%.1f%%" % (n1 / total * 100)),
                         ("第二次作业", total, n2, total - n2, "%.1f%%" % (n2 / total * 100))], start=1):
    r = s + j
    for i, v in enumerate(row):
        col = 1 if i == 0 else i + 2
        c = ws.cell(r, col, v)
        c.font = Font(name=FONT, size=10)
        c.alignment = CENTER
        c.border = BORDER
        c.fill = PatternFill("solid", fgColor="FFF2F6FA")
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
    ws.row_dimensions[r].height = 20

note = s + 4
c = ws.cell(note, 1, "说明：统计依据“作业”文件夹中各学生实际提交内容（代码 / 演示视频）核对；"
                     "交作业情况单元格中括号内文字为结构性缺失提示（缺源代码、缺演示视频、多次提交等），"
                     "不代表题目全部正确。")
c.font = Font(name=FONT, size=9, color="FF808080")
c.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
ws.merge_cells(start_row=note, start_column=1, end_row=note, end_column=8)
ws.row_dimensions[note].height = 34
c2 = ws.cell(note + 1, 1, "核对时间：2026-10-24")
c2.font = Font(name=FONT, size=9, color="FF808080")

wb.save(DST)
print("saved", DST, "last_row", last, "E", n1, "G", n2)
